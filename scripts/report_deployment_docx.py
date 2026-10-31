"""Build the Vietnamese deployment report from saved verification evidence."""
from pathlib import Path
import json
import re
from docx.image.image import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/docx/BAO_CAO_CICD_DOCKER_AI.docx'
E = ROOT / 'output/deployment'
if not (E / 'publication.json').exists():
    E = ROOT / 'deployment/verification/release-20261031'
def data(name):
    p = E / name
    return json.loads(p.read_text()) if p.exists() else {}
pub = data('publication.json')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin = Cm(3.5), Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)
for s in doc.styles:
    if s.type == 1:
        s.font.name = 'Times New Roman'
        s.font.size = Pt(14)
        s.font.color.rgb = RGBColor(0, 0, 0)
        fonts = s.element.get_or_add_rPr()
        rf = fonts.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); fonts.append(rf)
        for attr in list(rf.attrib):
            if 'Theme' in attr: del rf.attrib[attr]
        ppr = s.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
        for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            rf.set(qn('w:' + k), 'Times New Roman')
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.5
normal.space_after = Pt(6)
normal.first_line_indent = Cm(1)
normal.widow_control = True
for n, size in [('Title', 20), ('Heading 1', 16), ('Heading 2', 14)]:
    s = doc.styles[n]; s.font.size = Pt(size); s.font.bold = True
    f = s.paragraph_format; f.first_line_indent = Cm(0)
    f.space_before = Pt(12); f.space_after = Pt(8); f.keep_with_next = True
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.first_line_indent = Cm(0)
r = foot.add_run(); r.font.size = Pt(12)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); r._r.addnext(fld)

def para(t, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not style and 'https://' not in t and not re.search(r'[a-f0-9]{40}', t) else WD_ALIGN_PARAGRAPH.LEFT
    # Markdown is only used for lightweight emphasis in the source guide.
    for i, part in enumerate(re.split(r'(\*\*.*?\*\*)', t.replace('`', ''))):
        r = p.add_run(part[2:-2] if part.startswith('**') else part)
        r.bold = True if part.startswith('**') else None
    if t.endswith(':'): p.paragraph_format.keep_with_next = True
    if t.startswith('- '):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(-.35)
        p.paragraph_format.left_indent = Cm(.35)
    return p

def heading(t):
    return para(t.replace(':', '').replace('+', 'và').replace('/', ' '), 'Heading 1')
def code(lines):
    for index, line in enumerate(lines):
        for chunk in [line]:
            p = para(chunk)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = index < len(lines) - 1
            f = p.paragraph_format; f.first_line_indent = Cm(0); f.line_spacing = 1.0; f.space_after = Pt(2)
            for r in p.runs: r.font.size = Pt(11)

def table(rows):
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(6), Cm(9.5)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Cm(6), Cm(9.5)
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0: trpr.append(OxmlElement('w:tblHeader'))
        for c, text in zip(cells, row):
            p = c.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(5)
            r = p.add_run(text.replace('`', '')); r.font.size = Pt(12); r.bold = index == 0
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side); edge.set(qn('w:w'), '90'); edge.set(qn('w:type'), 'dxa'); margins.append(edge)
            pr.append(margins)
            align = OxmlElement('w:vAlign'); align.set(qn('w:val'), 'center'); pr.append(align)
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EEEEEE'); pr.append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

p = para('BÁO CÁO TRIỂN KHAI CI CD VÀ DOCKER CHO MÔ HÌNH AI', 'Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = para('Dự án CA TCC và DINOv3'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
para('Ngày thực hiện 31 tháng 10 năm 2026')
para('Báo cáo trình bày cách đóng gói hai mô hình nhận dạng âm thanh thành Docker image, tự động kiểm tra và phát hành bằng GitHub Actions, đồng thời đưa bản web lên GitHub Pages. Phần hướng dẫn đi từ chuẩn bị tài nguyên đến kiểm tra kết quả, kèm các địa chỉ và số liệu thực tế để đối chiếu.')
para('Phương án không cần thuê VPS. Bản web chạy suy luận ngay trong trình duyệt, nên người dùng vẫn có thể truy cập khi máy cá nhân của người triển khai đã tắt. Container Docker cung cấp API riêng; nếu chạy container trên máy cá nhân thì API đó chỉ hoạt động khi máy và Docker còn chạy.')
heading('1 Kết quả triển khai')
table([['Hạng mục', 'Kết quả đã ghi nhận'], ['CI', '15 kiểm tra logic, HTTP và tài nguyên đạt trên GitHub Actions.' if pub.get('ci_verified') else 'Chưa xác nhận.'], ['Docker Hub', 'Đã push, pull và kiểm tra API sau phát hành.' if pub.get('dockerhub_verified') else 'Đang chờ xác nhận phát hành.'], ['GitHub Pages', 'Đã triển khai và kiểm tra tài nguyên web công khai.' if pub.get('pages_verified') else 'Đang chờ xác nhận triển khai.'], ['Mô hình', 'RAVDESS gồm 8 cảm xúc, tối đa 6 giây; SpeechCommands gồm 35 từ khóa, tối đa 1 giây.']])
for label, key in [('Repository', 'repo_url'), ('CI', 'ci_url'), ('Phát hành', 'release_run_url'), ('Docker Hub', 'dockerhub_url'), ('Web', 'pages_url')]:
    if pub.get(key): para(label + ': ' + pub[key])
para('Commit mã nguồn dùng để phát hành: ' + pub.get('source_commit', 'chưa xác nhận'))
if pub.get('image_digest'): para('Digest image công bố: ' + pub['image_digest'])
heading('Sử dụng image đã phát hành')
para('Bản phát hành hiện tại dùng tag theo toàn bộ commit, không cần giả định đã có tag latest hoặc v1.0.0. Có thể tải và chạy đúng image bằng các lệnh sau.')
code(['docker pull ' + pub.get('image_tag', 'IMAGE_TAG'), 'docker run -d --name ca-tcc-public \\', '  -p 127.0.0.1:8000:8000 \\', '  ' + pub.get('image_digest', 'IMAGE_DIGEST'), 'curl --fail http://127.0.0.1:8000/healthz'])
para('Mở http://127.0.0.1:8000 để dùng giao diện của container. Nếu cổng 8000 đang có container thử nghiệm, dừng container đó trước khi chạy bản public.')
heading('2 Tổ chức hệ thống')
para('Nhánh Docker nhận WAV qua HTTP, tính đặc trưng log mel, chuẩn hóa và chạy TorchScript bằng CPU. Kết quả trả về dưới dạng JSON gồm phiên bản mô hình và các nhãn có điểm cao nhất. Image dùng Python 3.11 và PyTorch CPU 2.5.1; tiến trình chạy bằng UID 10001 và nạp sẵn hai mô hình trước khi nhận yêu cầu.')
para('Nhánh web phân phối HTML, JavaScript, WASM và trọng số ONNX. Trình duyệt tải trọng số của tác vụ được chọn rồi xử lý âm thanh trên thiết bị người dùng. Mỗi mô hình tải khoảng 84 MiB ở lần đầu, vì vậy thời gian khởi động phụ thuộc đường truyền và bộ nhớ thiết bị.')
para('CI chạy khi push vào main hoặc cập nhật pull request. Phát hành chạy thủ công hoặc bằng tag bắt đầu bằng v. Workflow lấy đúng bộ model đã ghim SHA256, kiểm tra, build image, thử API, push Docker Hub, pull kiểm tra lại rồi mới triển khai Pages. Quy trình này không tự huấn luyện lại mô hình.')

# Keep the reproducible guide, with report-specific introductory/result sections above.
lines = (ROOT / 'deployment/README.md').read_text().splitlines()
active = False; in_code = False; buf = []; rows = []; narrative = []
def flush():
    if narrative: para(' '.join(narrative)); narrative.clear()
def flush_table():
    if rows: table(rows); rows.clear()
for line in lines:
    if line.startswith('## 2.'): active = True
    if line.startswith('## 12.'): break
    if not active: continue
    if line.startswith('```'):
        flush(); flush_table()
        if in_code: code(buf); buf.clear()
        in_code = not in_code; continue
    if in_code: buf.append(line); continue
    if line.startswith('|'):
        flush()
        if not re.match(r'^\|[\s:|-]+\|$', line): rows.append([v.strip() for v in line.strip('|').split('|')])
        continue
    flush_table()
    if line.startswith('## '):
        flush()
        number, title = line[3:].split('. ', 1)
        heading(f'{int(number)+1} {title}'); continue
    if not line.strip(): flush(); continue
    if line.startswith('- ') or re.match(r'^\d+\. ', line):
        flush(); para(line if line.startswith('- ') else '- ' + re.sub(r'^\d+\. ', '', line)); continue
    if line.startswith(('Mọi URL Docker Hub', 'Không viết URL dự kiến', 'Ghi lại URL,')): continue
    narrative.append(line)
flush(); flush_table()
heading('13 Đối chiếu kết quả kiểm thử')
rows = [['Phép kiểm tra', 'Kết quả thực tế']]
for file, kind in [('cloud-published/published-smoke.json','API sau khi pull image công bố'), ('onnx-verification.json','ONNX trên CPU')]:
    obj = data(file)
    for dataset, item in obj.get('real_wav_parity', obj.get('datasets', {})).items():
        err = item.get('max_probability_error', item.get('maximum_probability_error'))
        rows.append([kind + ' ' + dataset, f'Sai khác điểm tối đa {err:.8g}; nhãn đứng đầu khớp tham chiếu.'])
table(rows)
para('Sai khác trên đo sự tương đương của kết quả suy luận với hai WAV validation, không phải accuracy hay F1 của toàn bộ mô hình. Kiểm tra API còn bao gồm WAV hỏng, dataset không hợp lệ, âm thanh vượt giới hạn và Origin không được phép.')
para('Kiểm tra ONNX bằng ONNX Runtime trên CPU xác nhận phần xuất mô hình giữ kết quả số học. Phép kiểm tra này không thay thế thao tác chọn tệp, chạy dự đoán và tải kết quả trong trình duyệt. Chỉ ghi nhận những thao tác trình duyệt đã được lưu bằng chứng.')
para(pub.get('browser_verification_note', 'Chưa có kiểm tra suy luận hoàn chỉnh trên bản Pages bằng giao diện trình duyệt. Cần kiểm tra bổ sung trên thiết bị khác với một phiên trình duyệt mới sau khi tắt máy cá nhân.'))
para('Bộ model được phát hành riêng tại https://github.com/VU-TRUONG-AI/CA-TCC-vs-DINOv3/releases/tag/model-assets-v1. Archive được chia thành 38 phần và có index chứa checksum. SHA256 toàn bộ archive là ' + pub.get('model_assets_sha256', 'chưa xác nhận') + '.')
para('Bộ tệp kiểm chứng gồm publication.json, container-verification.json, onnx-verification.json và các artifact của workflow phát hành. Không ghi token Docker Hub trong báo cáo hoặc repository.')
for name, caption in [('ci-success.jpg','CI trên GitHub Actions'), ('release-success.jpg','Workflow phát hành Docker và Pages'), ('dockerhub.jpg','Image trên Docker Hub'), ('pages.jpg','Dự đoán RAVDESS trên GitHub Pages'), ('pages-speech.jpg','Dự đoán SpeechCommands trên GitHub Pages')]:
    path = E / name
    if path.exists():
        doc.add_page_break(); heading(caption)
        p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
        im = Image.from_file(str(path)); w, h = im.px_width, im.px_height
        p.add_run().add_picture(str(path), width=Cm(min(15.5, 19.5 * w / h)))
        para('Ảnh chụp kết quả thực tế ngày 31 tháng 10 năm 2026.')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
